"""Construye el libro general desde su definición actual; no calcula ni aplica el método."""
from pathlib import Path
from datetime import datetime
from copy import copy
import argparse
import json
import re
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, GradientFill, Border, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.properties import CalcProperties
from openpyxl.xml.functions import fromstring

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'producto/metodo/oferta_entrega_y_economia/v0.1/modelo/definicion_modelo.json'
OUTPUT = SOURCE.parent / 'Go2Rev_modelo_economico_v0.2.xlsx'


def build(output=OUTPUT):
    definition = json.loads(SOURCE.read_text(encoding='utf-8'))
    styles = []
    for item in definition['styles']:
        fill = fromstring(item['fill'])
        fill_type = GradientFill if any('gradientFill' in e.tag for e in fill) else PatternFill
        styles.append({key: cls.from_tree(fromstring(item[key])) for key, cls in [('font',Font),('fill',fill_type),('border',Border),('alignment',Alignment),('protection',Protection)]} | {'number_format':item['number_format']})
    wb = Workbook(); wb.remove(wb.active)
    wb.properties.creator = 'RevOps Studio'
    wb.properties.title = 'Go2Rev Modelo económico v0.2'
    wb.properties.created = wb.properties.modified = datetime(2026,9,11)
    wb.calculation = CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    inputs = 0
    for spec in definition['sheets']:
        ws = wb.create_sheet(spec['name'])
        ws.sheet_view.showGridLines = spec['show_gridlines']
        ws.freeze_panes = spec['freeze']
        for a, props in spec['rows'].items():
            for key, value in props.items():
                if value is not None:setattr(ws.row_dimensions[int(a)],key,value)
        for a, props in spec['columns'].items():
            for key, value in props.items():
                if value is not None:setattr(ws.column_dimensions[a],key,value)
        for item in spec['cells']:
            if item.get('input'):
                inputs += 1
                if item.get('v') is not None:raise ValueError('Entrada empresarial en la definición general')
            if isinstance(item.get('v'),(int,float)):raise ValueError('Constante numérica ajena a las fórmulas generales')
            c=ws[item['a']];c.value=item.get('v')
            for key,value in styles[item['s']].items():setattr(c,key,copy(value))
        for area in spec['merges']:ws.merge_cells(area)
        for item in spec['validations']:ws.add_data_validation(DataValidation.from_tree(fromstring(item)))
        for item in spec['conditional_formats']:
            rule = Rule.from_tree(fromstring(item['xml']))
            rule.dxf = DifferentialStyle.from_tree(fromstring(item['dxf']))
            ws.conditional_formatting.add(item['range'],rule)
        ws.print_area=spec['print_area']
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=spec.get('print_height',0)
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.sheet_properties.tabColor='17324D'
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output)
    # Fechas de contenedor fijas; no altera fórmulas ni introduce resultados cacheados.
    with zipfile.ZipFile(output) as z:parts={name:z.read(name) for name in z.namelist()}
    parts['docProps/core.xml']=re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',rb'\g<1>2026-09-11T00:00:00Z\g<2>',parts['docProps/core.xml'])
    with zipfile.ZipFile(output,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for name,data in sorted(parts.items()):
            info=zipfile.ZipInfo(name,(2026,9,11,0,0,0));info.create_system=3;info.external_attr=0o100644<<16;info.compress_type=zipfile.ZIP_DEFLATED
            z.writestr(info,data)
    print(json.dumps({'archivo':str(output),'hojas':len(wb.worksheets),'entradas_vacias':inputs,'formulas_evaluadas':False},ensure_ascii=True))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--salida',type=Path,default=OUTPUT,help='Ruta del libro general que se genera o sustituye; no usar una copia de prestación')
    build(parser.parse_args().salida)
