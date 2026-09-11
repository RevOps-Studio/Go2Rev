from pathlib import Path
import re, json, hashlib, zipfile, sys, os
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
IS_GUIDE='--guia' in sys.argv
ARCH=ROOT/('producto/paquete_fundacional/v0.3' if IS_GUIDE else 'producto/arquitectura/v0.2')
CURRENT_SOURCE=ARCH/'LEEME.md'
OUT=Path(__file__).resolve().parent/'salida'
OUT.mkdir(parents=True,exist_ok=True)
NAVY='151B73'
d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.49)
normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(5 if IS_GUIDE else 6)
normal.paragraph_format.line_spacing=1.12
for key,size in [('Title',36),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    st=d.styles[key];st.font.name='Times New Roman';st.font.size=Pt(size)
    st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(12 if key!='Title' else 0)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.page_break_before=False
    for border in st._element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    for fonts in st._element.findall('.//'+qn('w:rFonts')):
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
d.styles['Footer'].font.name='Arial'
d.styles['Footer'].font.size=Pt(8)
d.styles['Footer'].font.color.rgb=RGBColor.from_string(NAVY)

def rich(p,text):
    text=text.replace('`','')
    for part in re.split(r'(\[[^\]]+\]\([^\)]+\)|\*\*.*?\*\*)',text):
        if not part:continue
        match=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if match:
            label,target=match.groups()
            if not re.match(r'^[a-z]+://',target):
                target=quote(os.path.relpath((CURRENT_SOURCE.parent/target).resolve(),ARCH).replace('\\','/'),safe='/#')
            rid=p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
            run=OxmlElement('w:r');rp=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),NAVY);rp.append(color);run.append(rp)
            t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)
            continue
        b=part.startswith('**') and part.endswith('**')
        r=p.add_run(part[2:-2] if b else part)
        if b:r.bold=True

def para(text,style='Normal'):
    p=d.add_paragraph(style=style);rich(p,text)
    if style=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

h=s.header.paragraphs[0]
h.text='Go2Rev Guía de implementación' if IS_GUIDE else 'Go2Rev Arquitectura metodológica'
h.paragraph_format.space_before=Pt(10)
for r in h.runs:r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(NAVY)
pr=h._p.get_or_add_pPr();bd=OxmlElement('w:pBdr');top=OxmlElement('w:top')
for key,value in [('val','single'),('sz','36'),('space','10'),('color',NAVY)]:top.set(qn('w:'+key),value)
bd.append(top);pr.append(bd)
f=s.footer.paragraphs[0]
f.add_run('Base fundacional — ')
for instruction,suffix in [('PAGE',' de '),('NUMPAGES','')]:
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instruction);f._p.append(field)
    f.add_run(suffix)
for r in f.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(NAVY)

def table(rows):
    n=len(rows[0]); widths={2:[1.8,4.7],3:[1.45,2.7,2.35],4:[1.3,1.45,1.8,1.95],5:[1.0,1.3,1.4,1.4,1.4]}[n]
    if rows[0][0]=='Carpeta':widths=[2.05,2.2,2.25]
    t=d.add_table(rows=1,cols=n);t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+key),value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        row=t.rows[0] if i==0 else t.add_row()
        rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if i==0:rp.append(OxmlElement('w:tblHeader'))
        for j,(c,txt) in enumerate(zip(row.cells,values)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY if i==0 else ('F4F5FA' if i%2==0 else 'FFFFFF'));cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge,value in [('top','70'),('bottom','70'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            p=c.paragraphs[0];rich(p,txt)
            if i==0:p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name='Times New Roman';r.font.size=Pt(10.5)
                if i==0:r.font.color.rgb=RGBColor(255,255,255);r.bold=False
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=.5

p=para('Go2Rev','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
if IS_GUIDE:
    d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for st in ['Heading 1','Heading 2','Heading 3']:d.styles[st].font.color.rgb=RGBColor(0,0,0)
    h.text='Go2Rev Guía de implementación'
    for r in h.runs:r.font.size=Pt(8)
    for border in h._p.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    para('Guía de implementación','Heading 1')
    para('Versión del conjunto 0.3 · 11 de septiembre de 2026')
    para('Para el consultor responsable de la primera prestación. Esta guía explica cómo encuadrar el servicio, investigar y recomendar, conectar el diseño, comprobar la preparación y transferir la continuidad. La base fundacional está construida y conserva su estado teórico; sus procedimientos se aplicarán en un encargo autorizado.')
    para('Los seis capítulos reúnen preparación del entorno, carpetas, glosario, guía, alcance y composición de entregables. Las instrucciones sustantivas, contratos, plantillas vacías y modelo económico se localizan mediante los enlaces del paquete. Las fuentes editables de esta vista permanecen junto a ella; los cambios se incorporan primero a esas fuentes.')
    chapters=[('../../metodo/operacion_conversacional/v0.2/09_uso_por_entorno.md','1 Preparar el entorno'),('../../metodo/operacion_conversacional/v0.2/08_carpetas_y_guardado.md','2 Carpetas y guardado'),('../../metodo/operacion_conversacional/v0.2/00_glosario.md','3 Glosario de lectura'),('01_guia_de_implementacion.md','4 Implementación de la prestación'),('02_alcance_y_variantes.md','5 Alcance y variantes'),('03_entregables_y_recepcion.md','6 Composición de entregables y recepción')]
else:
    para('Arquitectura metodológica','Heading 1')
    para('Base fundacional aceptada por Carlos Estrada')
    para('Versión de arquitectura 0.2 · Base aceptada el 10 de septiembre de 2026')
    para('La arquitectura conserva diecisiete nodos, contratos y correspondencia con los once pasos de origen. Su desarrollo sustantivo está integrado en el paquete fundacional 0.3. Las pruebas del método se sitúan después del cierre de construcción y requieren un ámbito autorizado; no se atribuye aquí aplicación previa.')
    para('Esta vista reúne los cuatro archivos operativos actuales de arquitectura, con integración editorial del 11 de septiembre de 2026. Las fuentes editables permanecen en producto/arquitectura/v0.2. La vista se regenera desde ellas y conserva el estado teórico del producto.')
    chapters=[('01_arquitectura_metodologica.md','1 Arquitectura y decisiones de producto'),('02_informacion_y_contratos.md','2 Información y contratos de conocimiento'),('03_nodos_y_entregables.md','3 Nodos y entregables'),('04_dependencias_y_gobernanza.md','4 Dependencias suficiencia y gobernanza')]
para('Contenido','Heading 2')
for _,title in chapters:para(title)
for index,(name,title) in enumerate(chapters):
    if index==0:d.add_page_break()
    CURRENT_SOURCE=ARCH/name
    para(title,'Heading 1')
    lines=(ARCH/name).read_text(encoding='utf-8').splitlines()[1:]
    i=0; code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if line.startswith('```'):
            code=False if code else (line[3:] or 'text')
            continue
        if code:
            if code!='mermaid':para(line)
            continue
        if not line or line.startswith('<a id='):continue
        if line.startswith('|'):
            rows=[]
            while True:
                cells=[c.strip() for c in line.strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',c) for c in cells):rows.append(cells)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            table(rows);continue
        if line.startswith('### '):para(line[4:],'Heading 3');continue
        if line.startswith('## '):para(line[3:],'Heading 2');continue
        if line.startswith('- '):line='• '+line[2:]
        line=line.replace('El diagrama orienta la lectura.','El mapa funcional orienta la lectura; la fuente editable incluye también un diagrama de navegación.')
        p=para(line)
        if line.startswith('Versión '):p.paragraph_format.keep_with_next=True
d.core_properties.title='Go2Rev Guía de implementación' if IS_GUIDE else 'Go2Rev Arquitectura metodológica'
d.core_properties.subject='Conjunto fundacional v0.3' if IS_GUIDE else 'Arquitectura v0.2'
d.core_properties.author='RevOps Studio'
d.core_properties.comments=''
dest=ARCH/('Go2Rev_guia_de_implementacion_v0.3.docx' if IS_GUIDE else 'Go2Rev_arquitectura_metodologica_v0.2.docx')
d.save(dest)
# El paquete se produce limpio: no conserva partes opacas de otros documentos.
with zipfile.ZipFile(dest) as z:parts={n:z.read(n) for n in z.namelist() if n!='docProps/thumbnail.jpeg'}
for n in ['[Content_Types].xml','_rels/.rels']:
    from lxml import etree
    root=etree.fromstring(parts[n])
    for child in list(root):
        if any('thumbnail' in v.lower() for v in child.attrib.values()):root.remove(child)
    parts[n]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in sorted(parts.items()):
        info=zipfile.ZipInfo(n,(2026,9,11,0,0,0));info.create_system=3;info.external_attr=0o100644<<16;info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,b)
report={'documento':dest.name,'fuentes':[{'archivo':name,'sha256':hashlib.sha256((ARCH/name).read_bytes()).hexdigest()} for name,_ in chapters],'tipo':'vista_derivada','partes_con_contenido_ajeno':False}
(OUT/('procedencia_guia.json' if IS_GUIDE else 'procedencia_documental.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(dest)
